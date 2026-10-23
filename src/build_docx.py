# -*- coding: utf-8 -*-
"""生成 Word 计算报告"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ---------- 全局字体 ----------
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing = 1.4
st.paragraph_format.space_after = Pt(4)

for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

def setfont(run, size=10.5, bold=False, cn="宋体", en="Times New Roman", color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if color: run.font.color.rgb = color

def title(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(p.add_run(txt), 17, True, "黑体", "Times New Roman"); 
    p.paragraph_format.space_after = Pt(2)
    return p

def subtitle(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(p.add_run(txt), 10.5, False, "楷体")
    p.paragraph_format.space_after = Pt(14)
    return p

def h1(txt):
    p = doc.add_paragraph()
    setfont(p.add_run(txt), 13, True, "黑体")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def h2(txt):
    p = doc.add_paragraph()
    setfont(p.add_run(txt), 11, True, "黑体")
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    return p

def para(txt, indent=True, size=10.5):
    p = doc.add_paragraph()
    setfont(p.add_run(txt), size)
    if indent: p.paragraph_format.first_line_indent = Pt(21)
    return p

def bullet(txt):
    p = doc.add_paragraph()
    setfont(p.add_run(txt), 10.5)
    p.paragraph_format.left_indent = Pt(21)
    return p

def caption(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(p.add_run(txt), 9.5, True, "黑体")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    return p

def note(txt):
    p = doc.add_paragraph()
    setfont(p.add_run(txt), 9, False, "楷体")
    p.paragraph_format.space_after = Pt(6)
    return p

def mktable(header, rows, widths=None, hl=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(header):
        hdr[i].text = ""
        pp = hdr[i].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        setfont(pp.add_run(htxt), 9.5, True, "黑体")
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bold = bool(hl and hl(r, i))
            setfont(pp.add_run(str(v)), 9.5, bold)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t
